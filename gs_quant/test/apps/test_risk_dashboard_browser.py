"""
Copyright 2026 Senanur Çetin.
Licensed under the Apache License, Version 2.0 (the "License");
you may not use this file except in compliance with the License.
You may obtain a copy of the License at

  http://www.apache.org/licenses/LICENSE-2.0

Unless required by applicable law or agreed to in writing,
software distributed under the License is distributed on an
"AS IS" BASIS, WITHOUT WARRANTIES OR CONDITIONS OF ANY
KIND, either express or implied.  See the License for the
specific language governing permissions and limitations
under the License.
"""

import os
import pathlib
import re
import socket
import threading
import time

import numpy as np
import openpyxl
import pytest

# The browser tests are skipped where Playwright or a browser is missing, unless CI says they must run
REQUIRED = os.environ.get('RISK_APP_REQUIRE_BROWSER') == '1'

pytest.importorskip('starlette')
pytest.importorskip('uvicorn')
if REQUIRED:
    from playwright import sync_api
else:
    sync_api = pytest.importorskip('playwright.sync_api')

import uvicorn  # noqa: E402

from gs_quant.apps.risk_dashboard.app import create_app  # noqa: E402
from gs_quant.apps.risk_dashboard.settings import Settings  # noqa: E402
from gs_quant.test.apps.test_risk_dashboard_marketdata import DATES, fx_service, prices_for  # noqa: E402


def _chromium():
    """A browser the environment already has, else None for Playwright's own"""
    path = pathlib.Path(os.environ.get('PLAYWRIGHT_BROWSERS_PATH', '')) / 'chromium'
    return str(path) if path.is_file() else None


@pytest.fixture(scope='module')
def server(tmp_path_factory):
    with socket.socket() as probe:
        probe.bind(('127.0.0.1', 0))
        port = probe.getsockname()[1]
    settings = Settings(port=port, database=tmp_path_factory.mktemp('data') / 'runs.db', market_data='yahoo')
    app = create_app(settings)
    quotes = {
        'THYAO.IS': ('TRY', prices_for(1), DATES),
        'GARAN.IS': ('TRY', prices_for(2), DATES),
        'ASELS.IS': ('TRY', prices_for(3), DATES),
        'AAPL': ('USD', prices_for(4), DATES),
        'USDTRY=X': ('TRY', 30 + np.arange(len(DATES)) * 0.05, DATES),
    }
    app.state.application.market = fx_service(quotes)  # no network: the provider answers from memory
    instance = uvicorn.Server(uvicorn.Config(app, port=port, log_level='warning'))
    thread = threading.Thread(target=instance.run, daemon=True)
    thread.start()
    deadline = time.time() + 10
    while not instance.started and time.time() < deadline:
        time.sleep(0.05)
    yield f'http://127.0.0.1:{port}'
    instance.should_exit = True
    thread.join(timeout=10)


@pytest.fixture(scope='module')
def browser():
    with sync_api.sync_playwright() as playwright:
        try:
            instance = playwright.chromium.launch(executable_path=_chromium(), args=['--no-sandbox'])
        except Exception as e:  # no browser installed
            if REQUIRED:
                raise
            pytest.skip(f'Chromium is not available: {e}')
        yield instance
        instance.close()


def test_portfolio_save_open_report_and_delete(server, browser, tmp_path):
    page = browser.new_page(accept_downloads=True)
    problems = []
    page.on('pageerror', lambda e: problems.append(str(e)))
    page.on('console', lambda m: problems.append(m.text) if m.type in ('error', 'warning') else None)

    page.goto(server)
    page.wait_for_selector('#results:not([hidden])')
    page.click('input[name=mode][value=portfolio]')
    page.click('#portfolio-sample')
    page.wait_for_selector('#portfolio-section:not([hidden])')
    assert page.locator('#table-assets tbody tr').count() == 4  # three assets and the portfolio
    assert page.locator('#table-stress tbody tr').count() == 3  # the worst 1, 5 and 20 periods
    assert page.locator('#table-stress thead th').count() == 7  # window, return, dates and three assets

    page.fill('#run-name', 'Browser test')
    page.click('#save-form button')
    page.wait_for_selector('#report:not([hidden])')
    assert 'Browser test' in page.inner_text('#history')

    with page.expect_download() as download:
        page.click('#report')
    report = tmp_path / 'report.html'
    download.value.save_as(report)

    with page.expect_download() as workbook_download:
        page.click('#xlsx')
    workbook_path = tmp_path / 'run.xlsx'
    workbook_download.value.save_as(workbook_path)
    assert workbook_download.value.suggested_filename.endswith('.xlsx')
    assert openpyxl.load_workbook(workbook_path).sheetnames[:2] == ['Summary', 'Backtest']
    assert 'Portfolio' in openpyxl.load_workbook(workbook_path).sheetnames

    page.reload()
    page.click('#history li button[data-action=open]')
    page.wait_for_selector('#portfolio-section:not([hidden])')

    saved = browser.new_page()
    saved.on('pageerror', lambda e: problems.append(str(e)))
    saved.on('console', lambda m: problems.append(m.text) if m.type in ('error', 'warning') else None)
    saved.goto(report.as_uri())
    saved.wait_for_selector('#results:not([hidden])')
    assert saved.is_visible('#portfolio-section') and not saved.is_visible('#controls-panel')
    assert 'Browser test' in saved.inner_text('#provenance')

    # printed, the report is the results only, on a light page, and Chromium can make a PDF of it
    saved.emulate_media(media='print', color_scheme='dark')
    assert (
        not saved.is_visible('#print') and not saved.is_visible('.actions') and saved.is_visible('#portfolio-section')
    )
    assert saved.evaluate("getComputedStyle(document.body).backgroundColor") == 'rgb(255, 255, 255)'
    assert saved.pdf(format='A4')[:5] == b'%PDF-'
    saved.emulate_media(media='screen', color_scheme='light')
    assert saved.is_visible('#print')

    saved.click('#lang')  # the report can be read in either language
    saved.locator('#glance-title').filter(has_text='Bir bakışta risk').wait_for()
    assert 'Kayıtlı rapor' in saved.inner_text('#provenance') or 'KAYITLI RAPOR' in saved.inner_text('#provenance')

    # a second analysis, with other weights, to compare against
    page.fill('#weight-0', '70')
    page.press('#weight-0', 'Tab')
    page.wait_for_selector('#results:not(.loading)')
    page.fill('#run-name', 'Browser test, 70% calm')
    page.click('#save-form button')
    page.locator('#history li').nth(1).wait_for()
    assert page.is_disabled('#compare')
    page.locator('#history input[data-compare]').nth(0).check()
    page.locator('#history input[data-compare]').nth(1).check()
    page.click('#compare')
    page.wait_for_selector('#compare-panel:not([hidden])')
    assert page.locator('#table-compare thead th').count() == 2
    assert page.locator('#table-compare tbody tr').count() >= 15
    page.click('#compare-close')
    assert not page.is_visible('#compare-panel')

    page.once('dialog', lambda dialog: dialog.accept())
    page.click('#history li button[data-action=delete]')
    page.locator('#history li').nth(1).wait_for(state='detached')

    assert problems == []  # no script error and no Content Security Policy violation, in the app or in the report


def test_loading_prices_by_symbol(server, browser):
    page = browser.new_page()
    problems = []
    page.on('pageerror', lambda e: problems.append(str(e)))
    page.on('console', lambda m: problems.append(m.text) if m.type in ('error', 'warning') else None)

    page.goto(server)
    page.wait_for_selector('#market-field:not([hidden])')
    page.wait_for_selector(
        '#results:not([hidden])'
    )  # the first analysis must be done: it clears the status when it ends
    assert 'Yahoo Finance' in page.inner_text('#market-hint')

    page.fill('#symbols', 'THYAO.IS, GARAN.IS')  # one symbol only in this mode
    page.click('#market-load')
    page.locator('#status.error').wait_for()
    assert 'one symbol' in page.inner_text('#status')

    page.fill('#symbols', 'THYAO.IS')
    page.select_option('#market-history', 'all')
    page.click('#market-load')
    page.locator('#provenance').filter(has_text='prices from Yahoo Finance').wait_for()
    provenance = page.inner_text('#provenance').lower()
    assert 'thyao.is' in provenance and 'market data' in provenance

    page.click('input[name=mode][value=portfolio]')
    page.fill('#symbols', 'THYAO.IS, GARAN.IS, ASELS.IS')
    page.click('#market-load')
    page.wait_for_selector('#portfolio-section:not([hidden])')
    assert page.locator('#table-assets tbody tr').count() == 4
    assert 'ASELS.IS' in page.inner_text('#table-assets')

    # assets quoted in different currencies are converted to one
    page.fill('#symbols', 'THYAO.IS, AAPL')
    page.click('#market-load')
    page.locator('#provenance').filter(has_text='THYAO.IS, AAPL in TRY').wait_for()
    assert 'USDTRY=X' in page.text_content('#notes') and 'different currencies' in page.text_content('#notes')
    page.select_option('#market-base', 'USD')
    page.click('#market-load')
    page.locator('#provenance').filter(has_text='in USD').wait_for()
    page.select_option('#market-base', '')

    page.fill('#symbols', 'NOPE.IS')
    page.click('#market-load')
    page.locator('#status.error').filter(has_text='at least two').wait_for()

    assert problems == []


def test_switching_to_turkish_and_back(server, browser):
    page = browser.new_context(locale='en-US').new_page()
    problems = []
    page.on('pageerror', lambda e: problems.append(str(e)))
    page.on('console', lambda m: problems.append(m.text) if m.type in ('error', 'warning') else None)

    page.goto(server)
    page.wait_for_selector('#results:not([hidden])')
    assert page.inner_text('h1') == 'Risk Analytics Dashboard' and page.get_attribute('html', 'lang') == 'en'
    english_headline = page.inner_text('#headline')
    assert '%' in english_headline and '.' in english_headline

    page.click('#lang')
    page.locator('h1').filter(has_text='Risk Analizi Panosu').wait_for()
    assert page.get_attribute('html', 'lang') == 'tr' and page.inner_text('#lang') == 'English'
    page.wait_for_selector('#results:not(.loading)')
    text = page.inner_text('body')
    assert 'Bir bakışta risk' in text and 'Model doğrulaması' in text and 'Kayıtlı analizler' in text
    assert (
        'Risk at a glance' not in text and '**' not in text and '{' not in text
    )  # nothing left untranslated or unfilled
    headline = page.inner_text('#headline')
    assert (
        '%' in headline and ',' in headline and re.search(r'%\d+,\d\d', headline)
    )  # a decimal comma, the percent sign first
    axis = page.locator('#chart-var svg text').evaluate_all('(nodes) => nodes.map((node) => node.textContent)')
    assert any(re.fullmatch(r'−?%\d+,\d', label) for label in axis)  # the chart axes follow the same convention
    assert [label for label in axis if re.search(r'\d\.\d', label)] == []
    assert 'Basel trafik ışığı' in page.inner_text('#table-tests')
    assert page.get_attribute('#run-name', 'placeholder') == 'Bu analize ad ver'
    assert page.title() == 'Risk Analizi Panosu'

    # the portfolio views are translated too, and the choice survives a reload
    page.click('input[name=mode][value=portfolio]')
    page.click('#portfolio-sample')
    page.wait_for_selector('#portfolio-section:not([hidden])')
    assert 'Çeşitlendirme oranı' in page.inner_text('#takeaway-portfolio')
    assert 'En kötü' in page.inner_text('#table-stress') and 'Sakin piyasa' in page.inner_text('#table-assets')
    page.reload()
    page.locator('h1').filter(has_text='Risk Analizi Panosu').wait_for()
    page.wait_for_selector('#results:not([hidden])')

    page.click('#lang')
    page.locator('h1').filter(has_text='Risk Analytics Dashboard').wait_for()
    assert page.get_attribute('html', 'lang') == 'en' and page.inner_text('#lang') == 'Türkçe'
    page.wait_for_selector('#results:not(.loading)')
    assert 'Risk at a glance' in page.inner_text('body')
    assert problems == []


def test_a_turkish_browser_gets_turkish_and_a_server_error_is_translated(server, browser):
    page = browser.new_context(locale='tr-TR').new_page()
    page.goto(server)
    page.wait_for_selector('#results:not([hidden])')
    assert page.inner_text('h1') == 'Risk Analizi Panosu' and page.inner_text('#lang') == 'English'

    page.click('input[name=mode][value=portfolio]')
    page.fill('#symbols', 'NOPE.IS, NOPE2.IS')
    page.click('#market-load')
    page.locator('#status.error').wait_for()
    assert 'bulunamadı' in page.inner_text('#status') or 'tanımıyor' in page.inner_text('#status')


def test_what_if_scenarios(server, browser):
    page = browser.new_context(locale='en-US').new_page()
    problems = []
    page.on('pageerror', lambda e: problems.append(str(e)))
    page.on('console', lambda m: problems.append(m.text) if m.type in ('error', 'warning') else None)
    page.goto(server)
    page.wait_for_selector('#results:not([hidden])')
    page.click('input[name=mode][value=portfolio]')
    page.click('#portfolio-sample')
    page.wait_for_selector('#portfolio-section:not([hidden])')
    assert page.is_visible('#whatif-add') and not page.is_visible('#table-whatif')

    # everything falls 10%: the portfolio falls 10%, a multiple of the VaR
    page.click('#whatif-add')
    effect = page.locator('#table-whatif .effect').first
    effect.filter(has_text='−10.00%').wait_for()
    assert re.fullmatch(r'\d+\.\d×', page.locator('#table-whatif .multiple').first.inner_text())
    assert page.locator('#table-whatif input[data-asset]').count() == 3

    # one asset stays put: the loss follows the weights (a third each)
    page.locator('#table-whatif input[data-asset]').first.fill('0')
    page.locator('#table-whatif input[data-asset]').first.press('Tab')
    effect.filter(has_text='−6.67%').wait_for()
    assert page.locator('#table-whatif input[data-asset]').first.input_value() == '0'  # what was typed is kept

    page.click('#whatif-add')
    page.locator('#table-whatif tbody tr').nth(1).wait_for()
    page.locator('#table-whatif input[data-field=name]').nth(1).fill('Second')
    page.locator('#table-whatif input[data-field=name]').nth(1).press('Tab')
    page.wait_for_selector('#results:not(.loading)')

    # saved with the analysis, and back when it is opened
    page.fill('#run-name', 'With scenarios')
    page.click('#save-form button')
    page.wait_for_selector('#report:not([hidden])')
    page.reload()
    page.wait_for_selector('#history li')
    page.click('#history li button[data-action=open]')
    page.locator('#table-whatif tbody tr').nth(1).wait_for()
    assert page.locator('#table-whatif input[data-field=name]').nth(1).input_value() == 'Second'
    page.locator('#table-whatif .effect').first.filter(has_text='−6.67%').wait_for()

    # removing a scenario, and the Turkish page
    page.locator('#table-whatif button[data-remove]').first.click()
    page.locator('#table-whatif tbody tr').nth(1).wait_for(state='detached')
    assert page.locator('#table-whatif input[data-field=name]').first.input_value() == 'Second'
    page.click('#lang')
    page.locator('#whatif-title').filter(has_text='Ya şöyle olursa').wait_for()
    assert 'Senaryo' in page.inner_text('#table-whatif') and 'Etki' in page.inner_text('#table-whatif')

    page.click('#lang')
    page.locator('#whatif-title').filter(has_text='What if?').wait_for()
    assert problems == []


def test_named_portfolios(server, browser):
    page = browser.new_context(locale='en-US').new_page()
    problems = []
    page.on('pageerror', lambda e: problems.append(str(e)))
    page.on('console', lambda m: problems.append(m.text) if m.type in ('error', 'warning') else None)
    page.goto(server)
    page.wait_for_selector('#results:not([hidden])')
    assert not page.is_visible('#book-panel')  # a series, not a portfolio
    page.click('input[name=mode][value=portfolio]')
    page.wait_for_selector('#book-panel:not([hidden])')
    assert not page.is_visible('#book-form')  # nothing from the market is loaded yet

    page.fill('#symbols', 'THYAO.IS, GARAN.IS')
    page.select_option('#market-base', 'USD')
    page.select_option('#market-history', '10')
    page.click('#market-load')
    page.wait_for_selector('#portfolio-section:not([hidden])')
    page.wait_for_selector('#book-form:not([hidden])')
    page.fill('#weight-0', '70')
    page.press('#weight-0', 'Tab')
    page.fill('#weight-1', '30')
    page.press('#weight-1', 'Tab')
    page.wait_for_selector('#results:not(.loading)')
    page.fill('#book-name', 'Banks')
    page.click('#book-form button')
    page.locator('#book li').first.wait_for()
    assert (
        'Banks' in page.inner_text('#book')
        and 'THYAO.IS, GARAN.IS' in page.inner_text('#book')
        and 'USD' in page.inner_text('#book')
    )

    # a new page: the portfolio is loaded again, with fresh prices and the weights that were saved
    page.reload()
    page.wait_for_selector('#results:not([hidden])')
    page.click('input[name=mode][value=portfolio]')
    page.locator('#book li').first.wait_for()
    page.click('#book li button[data-action=open]')
    page.wait_for_selector('#portfolio-section:not([hidden])')
    page.wait_for_selector('#results:not(.loading)')
    assert page.input_value('#symbols') == 'THYAO.IS, GARAN.IS' and page.input_value('#market-base') == 'USD'
    assert page.input_value('#market-history') == '10'
    assert page.input_value('#weight-0') == '70' and page.input_value('#weight-1') == '30'
    assert 'USD' in page.inner_text('#provenance')

    # saving under the same name replaces it
    page.fill('#weight-0', '50')
    page.press('#weight-0', 'Tab')
    page.fill('#weight-1', '50')
    page.press('#weight-1', 'Tab')
    page.wait_for_selector('#results:not(.loading)')
    page.click('#book-form button')
    page.locator('#status').filter(has_text='saved').wait_for()
    assert page.locator('#book li').count() == 1
    stored = page.request.get(f'{server}/api/portfolios').json()
    assert len(stored) == 1
    assert page.request.get(f'{server}/api/portfolios/{stored[0]["id"]}').json()['definition']['weights'] == {
        'THYAO.IS': 0.5,
        'GARAN.IS': 0.5,
    }

    page.click('#lang')
    page.locator('#book-title').filter(has_text='Portföylerim').wait_for()

    page.once('dialog', lambda dialog: dialog.accept())
    page.click('#book li button[data-action=delete]')
    page.locator('#book li').first.wait_for(state='detached')
    assert problems == []


def test_horizon_filtered_estimate_and_benchmark(server, browser):
    page = browser.new_context(locale='en-US').new_page()
    problems = []
    page.on('pageerror', lambda e: problems.append(str(e)))
    page.on('console', lambda m: problems.append(m.text) if m.type in ('error', 'warning') else None)
    page.goto(server)
    page.wait_for_selector('#results:not([hidden])')
    page.wait_for_selector('#results:not(.loading)')

    # the filtered estimate is always there, the benchmark only for a portfolio loaded with one
    assert page.locator('#table-ewma tbody tr').count() == 6
    assert page.locator('#headline > div').count() == 4
    assert not page.is_visible('#benchmark-section') and not page.is_visible('#benchmark-row')
    assert 'Filtered VaR (EWMA)' in page.inner_text('#chart-var')

    page.fill('#horizon', '10')
    page.press('#horizon', 'Tab')
    page.locator('#headline > div').nth(5).wait_for()
    assert '10 periods' in page.inner_text('#headline') and 'square root of time' in page.inner_text('#headline')
    assert 'square root of the horizon' in page.text_content('#notes')
    page.fill('#horizon', '1')
    page.press('#horizon', 'Tab')
    page.locator('#headline > div').nth(5).wait_for(state='detached')

    page.click('input[name=mode][value=portfolio]')
    page.wait_for_selector('#benchmark-row:not([hidden])')
    page.fill('#symbols', 'THYAO.IS, GARAN.IS')
    page.fill('#benchmark', 'aapl')
    page.click('#market-load')
    page.wait_for_selector('#benchmark-section:not([hidden])')
    page.wait_for_selector('#results:not(.loading)')
    assert page.locator('#table-assets tbody tr').count() == 3  # two assets and the portfolio: not the benchmark
    assert 'AAPL' in page.inner_text('#takeaway-benchmark')
    assert page.locator('#table-benchmark tbody tr').count() == 8
    assert 'Beta' in page.inner_text('#table-benchmark') and 'Tracking error' in page.inner_text('#table-benchmark')

    # a benchmark is kept with a named portfolio and comes back with it
    page.fill('#book-name', 'Banks vs Apple')
    page.click('#book-form button')
    page.locator('#book li').first.wait_for()
    stored = page.request.get(f'{server}/api/portfolios').json()
    assert [p['benchmark'] for p in stored if p['name'] == 'Banks vs Apple'] == ['AAPL']
    page.fill('#benchmark', '')
    page.click('#book li button[data-action=open]')
    sync_api.expect(page.locator('#benchmark')).to_have_value('AAPL')  # the saved definition is fetched first
    page.wait_for_selector('#benchmark-section:not([hidden])')
    page.wait_for_selector('#results:not(.loading)')

    page.click('#lang')
    page.locator('#benchmark-title').filter(has_text='Kıyasa göre').wait_for()
    assert 'zamanın karekökü' not in page.inner_text('#headline')  # a horizon of one shows no such figures
    assert 'Filtrelenmiş' in page.inner_text('#table-ewma')

    page.once('dialog', lambda dialog: dialog.accept())
    page.click('#book li button[data-action=delete]')
    page.locator('#book li').first.wait_for(state='detached')

    assert problems == []

    # a symbol that does not exist is an error, not a silent portfolio without a benchmark
    page.fill('#benchmark', 'NOPE.IS')
    page.click('#market-load')
    page.locator('#status.error').wait_for()
    assert problems == ['Failed to load resource: the server responded with a status of 422 (Unprocessable Entity)']


def test_the_static_export_works_in_a_browser(browser, tmp_path):
    """The exported dashboard (gs-quant-risk export) is one file that needs no server: it is what a demo site serves"""
    from gs_quant.apps.risk_dashboard import export

    page_file = tmp_path / 'dashboard.html'
    page_file.write_text(export.render(export.build_payload(n=400)), encoding='utf-8')
    page = browser.new_context(locale='en-US').new_page()
    problems = []
    page.on('pageerror', lambda e: problems.append(str(e)))
    page.on('console', lambda m: problems.append(m.text) if m.type in ('error', 'warning') else None)

    page.goto(page_file.as_uri())
    page.wait_for_selector('#results:not([hidden])')
    page.wait_for_selector('#results:not(.loading)')
    assert page.locator('#headline > div').count() == 4
    first = page.inner_text('#headline')
    for hidden in ('#mode-field', '#file-field', '#market-field', '#history-panel', '#book-panel', '#horizon-field'):
        assert not page.is_visible(hidden), hidden

    page.select_option('#scenario', 'regime_shift')
    page.select_option('#method', 'parametric')
    page.wait_for_selector('#results:not(.loading)')
    assert page.inner_text('#headline') != first
    assert page.locator('#chart-var svg').count() == 1 and page.locator('#table-ewma tbody tr').count() == 6
    # the slider steps through the precomputed levels: the last one is 99%
    page.eval_on_selector(
        '#confidence', "el => { el.value = 3; el.dispatchEvent(new Event('input', { bubbles: true })); }"
    )
    page.locator('#headline').filter(has_text='99.0% value at risk').wait_for()

    page.click('#lang')
    page.locator('#glance-title').filter(has_text='Bir bakışta risk').wait_for()
    assert problems == []
