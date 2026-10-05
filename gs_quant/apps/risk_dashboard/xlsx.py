"""
Copyright 2026 Senanur Çetin.
Licensed under the Apache License, Version 2.0 (the "License");
you may not use this file except in compliance with the License.
You may obtain a copy of the License at

  http://www.apache.org/licenses/LICENSE-2.0

Unless required by applicable law or agreed to in writing,
software distributed under the License is distributed on an
"AS IS" BASIS, WITHOUT WARRANTIES OR CONDITIONS OF ANY
KIND, either express or implied.  See the License for the
specific language governing permissions and limitations
under the License.
"""

import io
from typing import Any, Optional

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet

XLSX_MIME = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
PERCENT = '0.00%'
RATIO = '0.00'
BOLD = Font(bold=True)


def _put(sheet: Worksheet, row: int, column: int, value: Any, number_format: Optional[str] = None, bold=False) -> None:
    """Writes one cell. Text is always stored as text: a name that starts with = + - or @ must never become a formula."""
    cell = sheet.cell(row=row, column=column)
    if value is None:
        return
    if isinstance(value, str):
        cell.value = value
        cell.data_type = 's'
    elif isinstance(value, bool):
        cell.value = value
    else:
        cell.value = float(value) if isinstance(value, float) else value
        if number_format:
            cell.number_format = number_format
    if bold:
        cell.font = BOLD


def _table(sheet: Worksheet, start: int, heads: list[str], rows: list[list], formats: list[Optional[str]]) -> int:
    """A header row and the rows under it; returns the next free row"""
    for column, head in enumerate(heads, 1):
        _put(sheet, start, column, head, bold=True)
    for offset, values in enumerate(rows, 1):
        for column, value in enumerate(values, 1):
            _put(sheet, start + offset, column, value, formats[column - 1])
    return start + len(rows) + 2


def _pairs(sheet: Worksheet, start: int, heads: list[str], rows: list[tuple]) -> int:
    """Measure and value rows, each with the number format of its value; returns the next free row"""
    for column, head in enumerate(heads, 1):
        _put(sheet, start, column, head, bold=True)
    for offset, (label, value, number_format) in enumerate(rows, 1):
        _put(sheet, start + offset, 1, label)
        _put(sheet, start + offset, 2, value, number_format)
    return start + len(rows) + 2


def _widths(sheet: Worksheet, widths: list[int]) -> None:
    for column, width in enumerate(widths, 1):
        sheet.column_dimensions[sheet.cell(row=1, column=column).column_letter].width = width


def _summary(sheet: Worksheet, stored: dict, result: dict) -> None:
    s, settings = result['summary'], result['settings']
    _put(sheet, 1, 1, stored['name'], bold=True)
    _put(sheet, 2, 1, f'Saved {stored["created_at"]}. Results are fractions shown as percentages; losses are negative.')
    method = settings['method']
    es = 'historical' if method == 'historical' else 'parametric'
    rows = [
        ['Observations', s['observations'], None],
        ['Periods per year', settings['periods_per_year'], None],
        ['Annualized return', s['annualized_return'], PERCENT],
        ['Annualized volatility', s['annualized_volatility'], PERCENT],
        [f'Value at risk, one period ({method}, {settings["confidence"]:.1%})', s[f'var_{method}'], PERCENT],
        [f'Expected shortfall, one period ({es})', s[f'expected_shortfall_{es}'], PERCENT],
        [f'Value at risk, {result["horizon"]["periods"]} periods', result['horizon']['var'], PERCENT],
        [
            f'Expected shortfall, {result["horizon"]["periods"]} periods',
            result['horizon']['expected_shortfall'],
            PERCENT,
        ],
        ['Maximum drawdown', s['max_drawdown'], PERCENT],
        ['Worst period', s['worst_period'], PERCENT],
        ['Best period', s['best_period'], PERCENT],
        ['Sortino ratio', s['sortino_ratio'], RATIO],
        ['Calmar ratio', s['calmar_ratio'], RATIO],
        ['Omega ratio', s['omega_ratio'], RATIO],
        ['Ulcer index', s['ulcer_index'], PERCENT],
        ['Skewness', s['skewness'], RATIO],
        ['Excess kurtosis', s['excess_kurtosis'], RATIO],
        ['Downside deviation (annualized)', s['downside_deviation'], PERCENT],
    ]
    row = _pairs(sheet, 4, ['Measure', 'Value'], rows)
    e = result['ewma']
    window = settings['window']
    filtered = [
        ['Volatility, whole series (annualized)', e['full_volatility'], PERCENT],
        [f'Volatility, last {window} periods (annualized)', e['window_volatility'], PERCENT],
        [f'Volatility, filtered (EWMA, decay {e["decay"]})', e['volatility'], PERCENT],
        ['Value at risk, filtered (normal, one period)', e['var'], PERCENT],
        ['Breaches of the filtered value at risk', e['exceedances'], None],
        ['Periods judged', e['observations'], None],
        ['Kupiec p-value of the filtered value at risk', e['p_value'], RATIO],
    ]
    _put(sheet, row, 1, 'Rolling window against a filtered estimate', bold=True)
    row = _pairs(sheet, row + 1, ['Measure', 'Value'], filtered)
    _put(sheet, row, 1, 'Settings and notes', bold=True)
    for text in [
        f'Method: {method}; confidence {settings["confidence"]:.1%}; rolling window {window} periods',
        *result['assumptions'],
    ]:
        row += 1
        _put(sheet, row, 1, text)
    _widths(sheet, [56, 18])


def _backtest(sheet: Worksheet, result: dict) -> None:
    t, ind, light = result['backtest'], result['backtest']['independence'], result['backtest']['traffic_light']
    rows = [
        ['Kupiec (frequency)', t['lr_statistic'], t['p_value'], 'Rejected' if t['reject'] else 'Not rejected'],
        [
            'Christoffersen independence',
            ind['independence_lr'],
            ind['independence_p_value'],
            'Rejected' if ind['independence_reject'] else 'Not rejected',
        ],
        [
            'Conditional coverage',
            ind['conditional_coverage_lr'],
            ind['conditional_coverage_p_value'],
            'Rejected' if ind['conditional_coverage_reject'] else 'Not rejected',
        ],
    ]
    row = _table(sheet, 1, ['Test', 'LR statistic', 'p-value', 'Result'], rows, [None, '0.0000', '0.0000', None])
    counts = [
        ('Periods judged', t['observations'], None),
        ('Breaches', t['exceedances'], None),
        ('Expected breaches', t['expected_rate'] * t['observations'], '0.0'),
        ('Observed breach rate', t['observed_rate'], PERCENT),
        ('Basel traffic light', light['zone'], None),
        ('Cumulative probability', light['cumulative_probability'], '0.0000'),
    ]
    _pairs(sheet, row, ['Measure', 'Value'], counts)
    _widths(sheet, [32, 16, 12, 16])


def _series(sheet: Worksheet, result: dict) -> None:
    s = result['series']
    heads = ['Date', 'Return', 'Growth of 1', 'Drawdown', 'VaR', 'Expected shortfall', 'Filtered VaR (EWMA)', 'Breach']
    formats = [None, PERCENT, '0.0000', PERCENT, PERCENT, PERCENT, PERCENT, None]
    keys = ['dates', 'returns', 'growth', 'drawdown', 'var', 'expected_shortfall', 'ewma_var', 'breach']
    rows = [list(values) for values in zip(*(s[k] for k in keys))]
    _table(sheet, 1, heads, rows, formats)
    sheet.freeze_panes = 'A2'
    _widths(sheet, [12, 10, 12, 10, 10, 18, 20, 8])


def _stress(sheet: Worksheet, result: dict) -> None:
    rows = []
    names = [a['name'] for a in result['stress'][0].get('assets', [])] if result['stress'] else []
    for window in result['stress']:
        by_name = {a['name']: a['return'] for a in window.get('assets', [])}
        rows.append(
            [window['periods'], window['return'], window['start'], window['end'], *(by_name.get(n) for n in names)]
        )
    _table(
        sheet,
        1,
        ['Periods', 'Worst return', 'From', 'To', *names],
        rows,
        [None, PERCENT, None, None, *([PERCENT] * len(names))],
    )
    _widths(sheet, [10, 14, 12, 12])


def _portfolio(sheet: Worksheet, result: dict) -> None:
    p = result['portfolio']
    rows = [
        [
            a['name'],
            a['weight'],
            a['volatility'],
            a['var'],
            a['volatility_contribution'],
            a['es_contribution'],
        ]
        for a in p['assets']
    ]
    heads = ['Asset', 'Weight', 'Volatility', 'VaR alone', 'Share of volatility', 'Share of expected shortfall']
    row = _table(sheet, 1, heads, rows, [None, PERCENT, PERCENT, PERCENT, PERCENT, PERCENT])
    _put(sheet, row - 1, 1, 'Portfolio volatility (annualized)')
    _put(sheet, row - 1, 2, p['portfolio_volatility'], PERCENT)
    _put(sheet, row, 1, 'Diversification ratio')
    _put(sheet, row, 2, p['diversification_ratio'], RATIO)
    row += 2
    names = p['correlation']['names']
    _put(sheet, row, 1, 'Correlation of returns', bold=True)
    _table(
        sheet,
        row + 1,
        ['', *names],
        [[n, *p['correlation']['matrix'][i]] for i, n in enumerate(names)],
        [None] + [RATIO] * len(names),
    )
    _widths(sheet, [24, 12, 12, 12, 18, 26])


def _what_if(sheet: Worksheet, result: dict) -> None:
    rows = [[s['name'], s['loss'], s['var_multiple'], s['es_multiple']] for s in result['scenarios']]
    row = _table(
        sheet,
        1,
        ['Scenario', 'Effect on the portfolio', 'Times the VaR', 'Times the expected shortfall'],
        rows,
        [None, PERCENT, RATIO, RATIO],
    )
    for scenario in result['scenarios']:
        _put(sheet, row, 1, scenario['name'], bold=True)
        row = _table(
            sheet,
            row + 1,
            ['Asset', 'Weight', 'Shock', 'Contribution'],
            [[a['name'], a['weight'], a['shock'], a['contribution']] for a in scenario['assets']],
            [None, PERCENT, PERCENT, PERCENT],
        )
    _widths(sheet, [28, 22, 16, 28])


def _benchmark(sheet: Worksheet, result: dict) -> None:
    b = result['benchmark']
    rows = [
        ['Benchmark', b['name'], None],
        ['Periods in common', b['observations'], None],
        ['Beta', b['beta'], RATIO],
        ['Correlation', b['correlation'], RATIO],
        ['R-squared', b['r_squared'], PERCENT],
        ['Alpha (annualized)', b['alpha'], PERCENT],
        ['Active return (annualized)', b['active_return'], PERCENT],
        ['Tracking error (annualized)', b['tracking_error'], PERCENT],
        ['Information ratio', b['information_ratio'], RATIO],
        ['Volatility, portfolio (annualized)', b['volatility'], PERCENT],
        ['Volatility, benchmark (annualized)', b['benchmark_volatility'], PERCENT],
    ]
    _pairs(sheet, 1, ['Measure', 'Value'], rows)
    _widths(sheet, [36, 18])


def render_xlsx(stored: dict, result: dict) -> bytes:
    """A saved run as a workbook: the figures as numbers (not text), one sheet per part of the analysis"""
    book = Workbook()
    summary = book.active
    summary.title = 'Summary'
    _summary(summary, stored, result)
    _backtest(book.create_sheet('Backtest'), result)
    _series(book.create_sheet('Series'), result)
    _stress(book.create_sheet('Stress'), result)
    if 'portfolio' in result:
        _portfolio(book.create_sheet('Portfolio'), result)
    if result.get('scenarios'):
        _what_if(book.create_sheet('What-if'), result)
    if 'benchmark' in result:
        _benchmark(book.create_sheet('Benchmark'), result)
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()
